#!/usr/bin/env python3
"""Generate core/balance.gd from the Balance Lab sheet of the design workbook.

The workbook is the source of truth for tuning values. Change a value there,
then run this script so the game picks it up:

    python3 design/export_balance.py

Requires openpyxl (`python3 -m pip install --user openpyxl`).
Run it from the repo root, or pass the workbook and output paths explicitly:

    python3 design/export_balance.py path/to/workbook.xlsx path/to/balance.gd
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is not installed. Run: python3 -m pip install --user openpyxl")

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = REPO_ROOT / "design" / "leaf-me-alone-workbook.xlsx"
DEFAULT_OUTPUT = REPO_ROOT / "core" / "balance.gd"
SHEET_NAME = "Balance Lab"

# Units whose whole-number values should become ints instead of floats.
INT_UNITS = {"players", "lots", "credits", "sessions"}
# Units appended to the constant name so call sites read clearly.
NAME_SUFFIX_UNITS = {"hours": "HOURS", "minutes": "MINUTES", "minutes/lot": "MINUTES_PER_LOT"}


def constant_name(variable: str, unit: str) -> str:
    name = variable.replace("%", " percent ")
    name = re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_").upper()
    if re.match(r"^[0-9]", name):
        name = "V_" + name
    suffix = NAME_SUFFIX_UNITS.get(unit.strip().lower())
    if suffix and not name.endswith(suffix):
        name = f"{name}_{suffix}"
    return name


def unit_note(unit: str) -> str:
    unit = unit.strip()
    if "%" in unit:
        return f"{unit}; stored as a fraction where 0.25 means 25%"
    return unit


def format_value(value: float, unit: str) -> tuple[str, str]:
    if float(value).is_integer() and unit.strip().lower() in INT_UNITS:
        return "int", str(int(value))
    text = repr(float(value))
    return "float", text


def main() -> None:
    workbook_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    output_path = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUTPUT

    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        sys.exit(f"Sheet '{SHEET_NAME}' not found in {workbook_path}")
    ws = wb[SHEET_NAME]

    header: list[str] | None = None
    rows: list[dict[str, object]] = []
    for raw in ws.iter_rows(values_only=True):
        cells = ["" if c is None else c for c in raw]
        if header is None:
            if str(cells[0]).strip() == "ID":
                header = [str(c).strip() for c in cells]
            continue
        record = dict(zip(header, cells))
        row_id = str(record.get("ID", "")).strip()
        if not row_id.startswith("BAL-"):
            continue
        rows.append(record)

    if not rows:
        sys.exit("No BAL- rows found. Has the Balance Lab layout changed?")

    seen: set[str] = set()
    lines = [
        "# GENERATED FILE. Do not edit by hand.",
        f"# Source: design/{workbook_path.name}, sheet '{SHEET_NAME}'.",
        "# To change a value, edit the workbook and run: python3 design/export_balance.py",
        "class_name Balance",
        "extends RefCounted",
        "",
        "## Tuning values from the design workbook's Balance Lab. Every constant keeps its",
        "## BAL id so a value in code can be traced back to the workbook row it came from.",
        "## Most of these are working assumptions, not canon. Check the status in each comment.",
        "",
    ]
    for record in rows:
        row_id = str(record["ID"]).strip()
        variable = str(record.get("Variable", "")).strip()
        unit = str(record.get("Unit", "")).strip()
        value = record.get("Value")
        if not isinstance(value, (int, float)):
            sys.exit(f"{row_id} ({variable}) has a non-numeric value: {value!r}")
        name = constant_name(variable, unit)
        if name in seen:
            sys.exit(f"Duplicate constant name {name} from {row_id}; rename the variable in the workbook")
        seen.add(name)
        gd_type, gd_value = format_value(value, unit)
        classification = str(record.get("Classification", "")).strip()
        status = str(record.get("Status", "")).strip()
        rationale = str(record.get("Rationale", "")).strip()
        lines.append(f"## {row_id}: {variable} ({unit_note(unit)}). {classification}, status {status}.")
        if rationale:
            lines.append(f"## {rationale}")
        lines.append(f"const {name}: {gd_type} = {gd_value}")
        lines.append("")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8", newline="\n")
    print(f"Wrote {len(rows)} constants to {output_path}")


if __name__ == "__main__":
    main()
